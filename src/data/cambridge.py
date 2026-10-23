"""Read-only intake for the University of Cambridge carbon-fibre-tow friction data.

Source: University of Cambridge research repository, item CAM-1671,
"Friction of carbon fibre tows". The tower-on-tool workbook records a
measured coefficient of friction for a carbon-fibre tow sliding on a glass
tool under an increasing normal-load ramp, for two tow sizings.

This module only *reads* the preserved workbooks. It never rewrites them and
it never mixes the Cambridge carbon-fibre tow system with AI-TriboTwin's
original glass-fibre/epoxy composite records.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from src.config import PROJECT_ROOT

CAMBRIDGE_DIR = PROJECT_ROOT / "data" / "external" / "cambridge-cam-1671"
CAMBRIDGE_RAW_DIR = CAMBRIDGE_DIR / "raw"
CAMBRIDGE_PROCESSED_DIR = CAMBRIDGE_DIR / "processed"
TOW_TOOL_WORKBOOK = CAMBRIDGE_RAW_DIR / "TowTool_ContactLengthAndFrictionForce.xlsx"
TOW_TOOL_SHEET = "Sheet1"

# 1-based worksheet column indices used by the tow-on-tool workbook.
_COLUMNS = {
    "step": 1,
    "normal_load_n": 6,
    "contact_width_um": 7,
    "contact_width_mm": 8,
    "contact_area_mm2": 9,
    "nominal_pressure_mpa": 10,
    "nominal_pressure_kpa": 11,
    "contact_length_px": 12,
    "contact_length_scanned_mm": 13,
    "contact_length_lmic_mm": 14,
    "contact_half_width_a_mm": 15,
    "total_contact_area_mm2": 16,
    "contact_length_over_ideal": 21,
    "percent_contact_length": 22,
    "friction_force_n": 25,
    "friction_force_zeroed_n": 26,
    "friction_force_per_contact_n": 27,
    "coefficient_of_friction": 30,
    "friction_force_over_area_mpa": 32,
}

# Rows that carry the "Test N" label for each tow-type block, and the tow label.
_TOW_BLOCKS: tuple[tuple[str, tuple[int, ...]], ...] = (
    ("T700-12k-60E", (21, 39, 57, 75, 93)),
    ("T700-12k-50C", (119, 137, 155, 173, 191)),
)

# Features are deliberately restricted to the applied normal load and the
# normal-side contact measurements. The friction-force columns (and any
# column derived from them) are excluded so the coefficient of friction is
# never reconstructed from its own inputs.
MODEL_FEATURE_COLUMNS: tuple[str, ...] = (
    "normal_load_n",
    "log10_normal_load_n",
    "nominal_pressure_mpa",
    "contact_width_mm",
    "contact_area_mm2",
    "contact_length_scanned_mm",
    "contact_half_width_a_mm",
    "percent_contact_length",
    "tow_type",
)
LEAKAGE_COLUMNS: tuple[str, ...] = (
    "friction_force_n",
    "friction_force_zeroed_n",
    "friction_force_per_contact_n",
    "coefficient_of_friction",
    "friction_force_over_area_mpa",
)
TARGET_COLUMN = "coefficient_of_friction"
GROUP_COLUMN = "test_run_id"


class CambridgeDataError(RuntimeError):
    """Raised when the Cambridge workbook cannot be read as expected."""


@dataclass(frozen=True)
class CambridgeTowToolDataset:
    """Tidy tow-on-tool measurements plus traceable source provenance."""

    measurements: pd.DataFrame
    metadata: dict[str, object]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell(worksheet, row: int, column: int) -> object:
    return worksheet.cell(row=row, column=column).value
def _read_test_run(worksheet, tow_type: str, test_index: int, label_row: int) -> list[dict[str, object]]:
    """Read one load-ramp test block beginning at its 'Test N' label row."""
    header_row = label_row + 1
    header_cell = _cell(worksheet, header_row, 1)
    if header_cell is not None and str(header_cell).strip() != "Step":
        raise CambridgeDataError(
            f"Expected a 'Step' header at row {header_row} for {tow_type} test {test_index}."
        )
    run_id = f"{tow_type}-T{test_index}"
    records: list[dict[str, object]] = []
    row = label_row + 2
    while True:
        step = _cell(worksheet, row, 1)
        if not isinstance(step, (int, float)):
            break
        record: dict[str, object] = {
            "tow_type": tow_type,
            "test_index": test_index,
            "test_run_id": run_id,
            "source_row": row,
        }
        for name, column in _COLUMNS.items():
            record[name] = _cell(worksheet, row, column)
        records.append(record)
        row += 1
    if not records:
        raise CambridgeDataError(f"No step rows found for {tow_type} test {test_index}.")
    return records


def load_towtool_measurements(workbook_path: Path = TOW_TOOL_WORKBOOK) -> CambridgeTowToolDataset:
    """Parse the tow-on-tool friction workbook into a traceable tidy table."""
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Cambridge tow-on-tool workbook not found: {workbook_path}")

    # The workbook is small; normal (non read-only) mode is used because the
    # parser reads cells by coordinate, which is prohibitively slow when
    # openpyxl is in read-only mode.
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    if TOW_TOOL_SHEET not in workbook.sheetnames:
        raise CambridgeDataError(f"Workbook is missing the {TOW_TOOL_SHEET!r} sheet.")

    records: list[dict[str, object]] = []
    for tow_type, label_rows in _TOW_BLOCKS:
        for test_index, label_row in enumerate(label_rows, start=1):
            records.extend(_read_test_run(workbook[TOW_TOOL_SHEET], tow_type, test_index, label_row))
    workbook.close()

    frame = pd.DataFrame.from_records(records)
    frame["step"] = pd.to_numeric(frame["step"], errors="raise").astype(int)
    for name in _COLUMNS:
        if name == "step":
            continue
        frame[name] = pd.to_numeric(frame[name], errors="coerce")
    load = pd.to_numeric(frame["normal_load_n"], errors="coerce")
    frame["log10_normal_load_n"] = np.log10(load.where(load > 0).to_numpy(dtype=float))
    frame = frame.sort_values(["tow_type", "test_index", "step"], kind="stable").reset_index(drop=True)

    test_run_ids = sorted(frame[GROUP_COLUMN].unique().tolist())
    metadata: dict[str, object] = {
        "source_name": "Friction of carbon fibre tows (University of Cambridge)",
        "repository_item": "https://www.repository.cam.ac.uk/items/3039d85c-0261-46b9-b249-fec061329100",
        "workbook": workbook_path.name,
        "workbook_sha256": sha256_file(workbook_path),
        "worksheet": TOW_TOOL_SHEET,
        "material_domain": "Carbon-fibre tow (T700 12k, 60E and 50C sizing) sliding on a glass tool",
        "measurement_protocol": "Tow-on-tool sliding friction under a 14-step increasing normal-load ramp",
        "target": TARGET_COLUMN,
        "target_definition": (
            "Workbook 'Coefficient of Friction' = friction force per contact / total normal load"
        ),
        "observation_count": int(len(frame)),
        "test_run_count": len(test_run_ids),
        "test_run_ids": test_run_ids,
        "tow_types": sorted(frame["tow_type"].unique().tolist()),
        "feature_columns": list(MODEL_FEATURE_COLUMNS),
        "excluded_leakage_columns": list(LEAKAGE_COLUMNS),
        "caveats": [
            "Five sequential load sweeps are labelled Test 1-5 per tow type; the workbook "
            "does not expose a specimen serial number, so independence of the sweeps is "
            "assumed from the labels rather than proven.",
            "The first ramp steps sit near zero applied load, where the ratio response is "
            "unstable; those rows are retained and reported, not silently dropped.",
        ],
    }
    return CambridgeTowToolDataset(measurements=frame, metadata=metadata)


def write_processed_dataset(
    dataset: "CambridgeTowToolDataset",
    output_dir: Path = CAMBRIDGE_PROCESSED_DIR,
) -> tuple[Path, Path]:
    """Persist the tidy tow-on-tool table and its provenance sidecar."""
    output_dir.mkdir(parents=True, exist_ok=True)
    table_path = output_dir / "towtool_cof_measurements.csv"
    provenance_path = output_dir / "towtool_cof_provenance.json"
    dataset.measurements.to_csv(table_path, index=False, lineterminator="\n")
    payload = dict(dataset.metadata)
    payload["measurements_file"] = table_path.name
    payload["measurements_sha256"] = sha256_file(table_path)
    provenance_path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    return table_path, provenance_path


MIN_INDEPENDENT_TEST_RUNS = 10


def build_towtool_dataset(output_dir: Path = CAMBRIDGE_PROCESSED_DIR) -> tuple[Path, Path]:
    """Convenience entry point: parse, validate, and persist the tow-on-tool dataset."""
    dataset = load_towtool_measurements()
    _validate_dataset(dataset.measurements, min_groups=MIN_INDEPENDENT_TEST_RUNS)
    return write_processed_dataset(dataset, output_dir)


def _validate_dataset(frame: pd.DataFrame, *, min_groups: int = 1) -> None:
    """Fail closed if the parsed table violates expected structural invariants."""
    required = {GROUP_COLUMN, TARGET_COLUMN, "tow_type", "test_index", "step", *MODEL_FEATURE_COLUMNS}
    missing = required - set(frame.columns)
    if missing:
        raise CambridgeDataError(f"Parsed dataset is missing columns: {sorted(missing)}")
    if frame.empty:
        raise CambridgeDataError("Parsed dataset is empty.")
    groups = frame[GROUP_COLUMN].nunique()
    if groups < min_groups:
        raise CambridgeDataError(
            f"Expected at least {min_groups} independent test runs; parsed {groups}."
        )
    if frame[GROUP_COLUMN].isna().any() or (frame[GROUP_COLUMN].astype(str).str.strip() == "").any():
        raise CambridgeDataError("Every observation requires a test-run identifier.")
    if not np.isfinite(frame[TARGET_COLUMN].to_numpy(dtype=float)).all():
        raise CambridgeDataError("Coefficient-of-friction column contains non-finite values.")
    duplicated = frame[["tow_type", "test_index", "step"]].duplicated().any()
    if duplicated:
        raise CambridgeDataError("Duplicate (tow_type, test_index, step) keys detected.")
